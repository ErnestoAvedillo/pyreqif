import io

from openpyxl import load_workbook
from PIL import Image as PILImage

from pyreqif import Reqif, ReqifError


def test_parses_requirements_and_chapters(sample_reqif_path):
    doc = Reqif(sample_reqif_path)

    assert len(doc) == 3
    chapter = doc.get('_chapter-1')
    assert chapter.is_chapter is True
    assert chapter.text == 'Introduccion'

    req1 = doc.get('_req-1')
    assert 'arrancar en menos de 2 segundos' in req1.text
    assert req1.images == ['attachments/img1.png']

    req2 = doc.get('_req-2')
    assert req2.comment == 'Ya validado en pruebas de carga.'
    assert req2.status == 'akzeptiert'


def test_status_options_come_from_the_enumeration_datatype(sample_reqif_path):
    doc = Reqif(sample_reqif_path)
    assert doc.status_options == ['akzeptiert', 'Klaerungsbedarf', 'nicht akzeptiert']


def test_update_unknown_requirement_raises(sample_reqif_path):
    doc = Reqif(sample_reqif_path)
    try:
        doc.update('_does-not-exist', comment='x')
        assert False, 'expected ReqifError'
    except ReqifError:
        pass


def test_update_and_save_roundtrip(sample_reqif_path, tmp_path):
    doc = Reqif(sample_reqif_path)
    doc.update('_req-1', comment='Comentario nuevo äöü', status='Klaerungsbedarf')

    out_path = tmp_path / 'edited.reqif'
    doc.save(out_path)

    reloaded = Reqif(out_path)
    req1 = reloaded.get('_req-1')
    assert req1.comment == 'Comentario nuevo äöü'
    assert req1.status == 'Klaerungsbedarf'
    # el texto original no se ha tocado
    assert 'arrancar en menos de 2 segundos' in req1.text


def test_update_none_leaves_field_untouched(sample_reqif_path):
    doc = Reqif(sample_reqif_path)
    doc.update('_req-2', comment=None, status='nicht akzeptiert')

    req2 = doc.get('_req-2')
    assert req2.comment == 'Ya validado en pruebas de carga.'
    assert req2.status == 'nicht akzeptiert'


def test_update_empty_string_clears_field(sample_reqif_path):
    doc = Reqif(sample_reqif_path)
    doc.update('_req-2', status='')

    req2 = doc.get('_req-2')
    assert req2.status == ''


def test_update_many_returns_count_of_matched_rows(sample_reqif_path):
    doc = Reqif(sample_reqif_path)
    applied = doc.update_many([
        ('_req-1', 'c1', 'akzeptiert'),
        ('_does-not-exist', 'c2', 'akzeptiert'),
    ])
    assert applied == 1
    assert doc.get('_req-1').comment == 'c1'


def test_to_excel_has_data_validation_with_status_options(sample_reqif_path, tmp_path):
    doc = Reqif(sample_reqif_path)
    xlsx_path = tmp_path / 'out.xlsx'
    doc.to_excel(xlsx_path)

    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == ['Requisitos', 'Opciones']

    sheet = workbook['Requisitos']
    assert [c.value for c in sheet[1]] == ['ID', 'Texto', 'Kommentar Lieferant M', 'Status Lieferant M']
    assert sheet.max_row == len(doc) + 1

    options = [c[0].value for c in workbook['Opciones'].iter_rows(min_row=2)]
    assert options == doc.status_options

    validations = list(sheet.data_validations.dataValidation)
    assert len(validations) == 1
    assert 'Opciones' in validations[0].formula1


def test_update_from_excel_applies_only_comment_and_status(sample_reqif_path, tmp_path):
    doc = Reqif(sample_reqif_path)
    xlsx_path = tmp_path / 'roundtrip.xlsx'
    doc.to_excel(xlsx_path)

    workbook = load_workbook(xlsx_path)
    sheet = workbook['Requisitos']
    # fila del _req-1: columna A=ID, C=comentario, D=estado
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == '_req-1':
            row[2].value = 'Actualizado desde Excel'
            row[3].value = 'nicht akzeptiert'
    workbook.save(xlsx_path)

    updated = doc.update_from_excel(xlsx_path)
    assert updated == len(doc)

    req1 = doc.get('_req-1')
    assert req1.comment == 'Actualizado desde Excel'
    assert req1.status == 'nicht akzeptiert'
    assert 'arrancar en menos de 2 segundos' in req1.text


def _make_png_bytes(size=(300, 150), color=(10, 200, 10)):
    buffer = io.BytesIO()
    PILImage.new('RGB', size, color=color).save(buffer, format='PNG')
    buffer.seek(0)
    return buffer


def test_to_excel_without_resolver_has_no_image_column(sample_reqif_path, tmp_path):
    doc = Reqif(sample_reqif_path)
    xlsx_path = tmp_path / 'no_images.xlsx'
    doc.to_excel(xlsx_path)

    sheet = load_workbook(xlsx_path)['Requisitos']
    assert [c.value for c in sheet[1]] == ['ID', 'Texto', 'Kommentar Lieferant M', 'Status Lieferant M']
    assert sheet._images == []


def test_to_excel_embeds_image_when_resolver_given(sample_reqif_path, tmp_path):
    doc = Reqif(sample_reqif_path)
    xlsx_path = tmp_path / 'with_images.xlsx'

    def resolver(image_ref):
        assert image_ref == 'attachments/img1.png'
        return _make_png_bytes()

    doc.to_excel(xlsx_path, image_resolver=resolver)

    sheet = load_workbook(xlsx_path)['Requisitos']
    assert [c.value for c in sheet[1]] == ['ID', 'Texto', 'Kommentar Lieferant M', 'Status Lieferant M', 'Imagen']
    # solo _req-1 tiene imagen, en la fila 3 (cabecera=1, _chapter-1=2, _req-1=3)
    assert len(sheet._images) == 1
    anchor = sheet._images[0].anchor
    assert anchor._from.col == 4  # columna E (0-indexed)
    assert anchor._from.row == 2  # fila 3 (0-indexed)


def test_to_excel_skips_missing_image_without_crashing(sample_reqif_path, tmp_path):
    doc = Reqif(sample_reqif_path)
    xlsx_path = tmp_path / 'missing_image.xlsx'

    doc.to_excel(xlsx_path, image_resolver=lambda ref: None)

    sheet = load_workbook(xlsx_path)['Requisitos']
    assert sheet._images == []


def test_to_excel_scales_down_large_images(sample_reqif_path, tmp_path):
    doc = Reqif(sample_reqif_path)
    xlsx_path = tmp_path / 'big_image.xlsx'

    doc.to_excel(xlsx_path, image_resolver=lambda ref: _make_png_bytes(size=(1000, 500)))

    sheet = load_workbook(xlsx_path)['Requisitos']
    image = sheet._images[0]
    assert image.width <= 160
    assert image.height <= 160
