"""Reqif: reading, editing, and exporting a single .reqif document."""
from __future__ import annotations

import io
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image as PILImage

XHTML_NS = 'http://www.w3.org/1999/xhtml'

EXCEL_HEADERS = ['ID', 'Texto', 'Kommentar Lieferant M', 'Status Lieferant M']
IMAGE_COLUMN = 'E'
IMAGE_MAX_SIZE = 160  # px, both width and height

_DEFINITION_TAGS = [
    'ATTRIBUTE-DEFINITION-STRING',
    'ATTRIBUTE-DEFINITION-XHTML',
    'ATTRIBUTE-DEFINITION-ENUMERATION',
    'ATTRIBUTE-DEFINITION-INTEGER',
    'ATTRIBUTE-DEFINITION-BOOLEAN',
    'ATTRIBUTE-DEFINITION-DATE',
]


class ReqifError(Exception):
    pass


@dataclass
class Requirement:
    """A ReqIF SPEC-OBJECT, with the relevant fields already resolved."""
    id: str
    text: str
    is_chapter: bool
    comment: str
    status: str
    images: list[str] = field(default_factory=list)


class Reqif:
    """Represents a single .reqif file (XML, OMG ReqIF standard).

    The comment and status of each requirement can be edited
    (`Kommentar Lieferant M` / `Status Lieferant M` by default, but the
    attribute names are configurable via subclassing or the
    TEXT_ATTR/CHAPTER_ATTR/COMMENT_ATTR/STATUS_ATTR class parameters),
    exported to Excel with data validation for the status, and imported
    back from a modified Excel file.
    """

    TEXT_ATTR = 'ReqIF.Text'
    CHAPTER_ATTR = 'ReqIF.ChapterName'
    COMMENT_ATTR = 'Kommentar Lieferant M'
    STATUS_ATTR = 'Status Lieferant M'

    def __init__(self, source):
        """source: path to a .reqif, or an already-open file-like object."""
        if hasattr(source, 'read'):
            self.path = None
            self._tree = ET.parse(source)
        else:
            self.path = Path(source)
            self._tree = ET.parse(self.path)

        self._root = self._tree.getroot()
        self._ns, self._uri = self._namespace()
        self._definitions = self._build_definitions()
        self._status_options = self._status_options_list()
        self._requirements, self._by_id = self._parse_requirements()

    # -- construction ----------------------------------------------------

    def _namespace(self):
        tag = self._root.tag
        if '}' in tag:
            uri = tag.split('}')[0].strip('{')
            return {'r': uri}, uri
        return {}, None

    def _qn(self, tag):
        return f'{{{self._uri}}}{tag}' if self._uri else tag

    def _build_definitions(self):
        ns = self._ns
        root = self._root
        enum_values_by_datatype = {}
        for dt in root.findall('.//r:DATATYPE-DEFINITION-ENUMERATION', ns):
            values = [
                (ev.attrib.get('IDENTIFIER'), ev.attrib.get('LONG-NAME', ''))
                for ev in dt.findall('.//r:ENUM-VALUE', ns)
            ]
            enum_values_by_datatype[dt.attrib.get('IDENTIFIER')] = values

        definitions = {}
        for tag in _DEFINITION_TAGS:
            for el in root.findall(f'.//r:{tag}', ns):
                ident = el.attrib.get('IDENTIFIER')
                entry = {'long_name': el.attrib.get('LONG-NAME', ''), 'options': []}
                if tag == 'ATTRIBUTE-DEFINITION-ENUMERATION':
                    ref = el.find('.//r:DATATYPE-DEFINITION-ENUMERATION-REF', ns)
                    if ref is not None:
                        entry['options'] = enum_values_by_datatype.get(ref.text, [])
                definitions[ident] = entry
        return definitions

    def _ids_by_name(self, long_name):
        return {ident for ident, meta in self._definitions.items() if meta['long_name'] == long_name}

    def _status_options_list(self):
        for ident in self._ids_by_name(self.STATUS_ATTR):
            return self._definitions[ident]['options']
        return []

    @staticmethod
    def _xhtml_text(the_value_el):
        if the_value_el is None:
            return ''
        return ''.join(the_value_el.itertext()).strip()

    @staticmethod
    def _collect_images(the_value_el):
        """<object type="image/..." data="path"> inside the XHTML,
        including nested ones (e.g. a .doc with a PNG preview
        as fallback)."""
        if the_value_el is None:
            return []
        images = []
        for obj in the_value_el.iter(f'{{{XHTML_NS}}}object'):
            obj_type = obj.attrib.get('type', '')
            data = obj.attrib.get('data')
            if data and obj_type.startswith('image/'):
                images.append(data)
        return images

    def _parse_requirements(self):
        ns = self._ns
        text_ids = self._ids_by_name(self.TEXT_ATTR)
        chapter_ids = self._ids_by_name(self.CHAPTER_ATTR)
        comment_ids = self._ids_by_name(self.COMMENT_ATTR)
        status_ids = self._ids_by_name(self.STATUS_ATTR)

        requirements = []
        by_id = {}
        for spec_object in self._root.findall('.//r:SPEC-OBJECT', ns):
            identifier = spec_object.attrib.get('IDENTIFIER', '')
            values = spec_object.find('r:VALUES', ns)
            if values is None:
                continue

            text = ''
            chapter = ''
            comment = ''
            status = ''
            images = []

            for xhtml_val in values.findall('r:ATTRIBUTE-VALUE-XHTML', ns):
                ref = xhtml_val.find('.//r:ATTRIBUTE-DEFINITION-XHTML-REF', ns)
                if ref is None or not ref.text:
                    continue
                the_value = xhtml_val.find('r:THE-VALUE', ns)
                plain = self._xhtml_text(the_value)
                images.extend(self._collect_images(the_value))
                if ref.text in text_ids:
                    text = plain
                elif ref.text in chapter_ids:
                    chapter = plain
                elif ref.text in comment_ids:
                    comment = plain

            for enum_val in values.findall('r:ATTRIBUTE-VALUE-ENUMERATION', ns):
                ref = enum_val.find('.//r:ATTRIBUTE-DEFINITION-ENUMERATION-REF', ns)
                if ref is None or ref.text not in status_ids:
                    continue
                enum_ref = enum_val.find('.//r:ENUM-VALUE-REF', ns)
                if enum_ref is not None:
                    for opt_id, opt_name in self._status_options:
                        if opt_id == enum_ref.text:
                            status = opt_name
                            break

            requirement = Requirement(
                id=identifier,
                text=text or chapter,
                is_chapter=bool(chapter and not text),
                comment=comment,
                status=status,
                images=images,
            )
            requirements.append(requirement)
            by_id[identifier] = requirement

        return requirements, by_id

    # -- public API ---------------------------------------------------

    @property
    def requirements(self) -> list[Requirement]:
        return self._requirements

    @property
    def status_options(self) -> list[str]:
        return [name for _opt_id, name in self._status_options]

    def get(self, identifier: str) -> Requirement | None:
        return self._by_id.get(identifier)

    def __len__(self):
        return len(self._requirements)

    def __iter__(self):
        return iter(self._requirements)

    def update(self, identifier: str, comment: str | None = None, status: str | None = None):
        """Updates the comment and/or status of a requirement.

        `None` leaves the field as is; use '' to clear it.
        """
        requirement = self._by_id.get(identifier)
        if requirement is None:
            raise ReqifError(f'Requirement {identifier} not found')

        spec_object = self._find_spec_object(identifier)
        values = spec_object.find('r:VALUES', self._ns)
        if values is None:
            values = ET.SubElement(spec_object, self._qn('VALUES'))

        if comment is not None:
            comment_def_id = next(iter(self._ids_by_name(self.COMMENT_ATTR)), None)
            if comment_def_id:
                self._set_xhtml_value(values, comment_def_id, comment)
            requirement.comment = comment.strip()

        if status is not None:
            status_def_id = next(iter(self._ids_by_name(self.STATUS_ATTR)), None)
            if status_def_id:
                status_ids_by_name = {name: opt_id for opt_id, name in self._status_options}
                self._set_enumeration_value(values, status_def_id, status, status_ids_by_name)
            requirement.status = status.strip()

    def update_many(self, updates: Iterable[tuple[str, str | None, str | None]]) -> int:
        """updates: iterable of (identifier, comment, status). Returns
        how many rows matched an existing requirement."""
        applied = 0
        for identifier, comment, status in updates:
            if identifier not in self._by_id:
                continue
            self.update(identifier, comment=comment, status=status)
            applied += 1
        return applied

    def _find_spec_object(self, identifier):
        for candidate in self._root.findall('.//r:SPEC-OBJECT', self._ns):
            if candidate.attrib.get('IDENTIFIER') == identifier:
                return candidate
        raise ReqifError(f'SPEC-OBJECT {identifier} not found')

    def _set_xhtml_value(self, values, def_id, text):
        ns = self._ns
        target = None
        for xhtml_val in values.findall('r:ATTRIBUTE-VALUE-XHTML', ns):
            ref = xhtml_val.find('.//r:ATTRIBUTE-DEFINITION-XHTML-REF', ns)
            if ref is not None and ref.text == def_id:
                target = xhtml_val
                break

        text = (text or '').strip()

        if target is None:
            if not text:
                return
            target = ET.SubElement(values, self._qn('ATTRIBUTE-VALUE-XHTML'))
            definition = ET.SubElement(target, self._qn('DEFINITION'))
            ref_el = ET.SubElement(definition, self._qn('ATTRIBUTE-DEFINITION-XHTML-REF'))
            ref_el.text = def_id

        old_value = target.find('r:THE-VALUE', ns)
        if old_value is not None:
            target.remove(old_value)

        the_value = ET.SubElement(target, self._qn('THE-VALUE'))
        div = ET.SubElement(the_value, f'{{{XHTML_NS}}}div')
        div.text = text

    def _set_enumeration_value(self, values, def_id, status_name, status_ids_by_name):
        ns = self._ns
        target = None
        for enum_val in values.findall('r:ATTRIBUTE-VALUE-ENUMERATION', ns):
            ref = enum_val.find('.//r:ATTRIBUTE-DEFINITION-ENUMERATION-REF', ns)
            if ref is not None and ref.text == def_id:
                target = enum_val
                break

        enum_id = status_ids_by_name.get((status_name or '').strip())

        if not enum_id:
            if target is not None:
                values.remove(target)
            return

        if target is None:
            target = ET.SubElement(values, self._qn('ATTRIBUTE-VALUE-ENUMERATION'))
            definition = ET.SubElement(target, self._qn('DEFINITION'))
            ref_el = ET.SubElement(definition, self._qn('ATTRIBUTE-DEFINITION-ENUMERATION-REF'))
            ref_el.text = def_id

        old_values = target.find('r:VALUES', ns)
        if old_values is not None:
            target.remove(old_values)

        values_el = ET.SubElement(target, self._qn('VALUES'))
        ref_val = ET.SubElement(values_el, self._qn('ENUM-VALUE-REF'))
        ref_val.text = enum_id

    # -- persistence ----------------------------------------------------

    def save(self, destination=None):
        """Writes the current XML to `destination` (path or file-like
        object), or to self.path if none is given."""
        target = destination if destination is not None else self.path
        if target is None:
            raise ReqifError('No destination path: pass one to save() or create the Reqif from a path.')
        self._tree.write(target, encoding='utf-8', xml_declaration=True)

    # -- Excel ----------------------------------------------------------

    @staticmethod
    def _scaled_png(source, max_size=IMAGE_MAX_SIZE):
        """Opens an image (path or file-like object), shrinks it if
        needed to fit within max_size x max_size px, and returns it
        as (BytesIO in PNG, (width, height))."""
        image = PILImage.open(source)
        image.load()
        if image.mode not in ('RGB', 'RGBA'):
            image = image.convert('RGBA')

        width, height = image.size
        scale = min(max_size / width, max_size / height, 1.0)
        new_size = (max(1, round(width * scale)), max(1, round(height * scale)))
        if new_size != (width, height):
            image = image.resize(new_size)

        buffer = io.BytesIO()
        image.save(buffer, format='PNG')
        buffer.seek(0)
        return buffer, new_size

    def to_excel(self, destination, image_resolver: Callable[[str], object] | None = None):
        """Writes to `destination` (path or file-like object) an
        .xlsx with one row per requirement and a validation dropdown
        in the status column with the actual ReqIF options.

        If `image_resolver` is given (a function that receives the path
        as it appears in `Requirement.images` and returns a path or a
        file-like object readable by PIL, e.g. `Reqifz.image_path`),
        the first image of each requirement is embedded in the cell of
        the "Imagen" column. An image that can't be read is skipped
        without interrupting the export.
        """
        option_names = [name for name in self.status_options if name]
        wrap_top = Alignment(wrap_text=True, vertical='top')
        headers = list(EXCEL_HEADERS) + (['Imagen'] if image_resolver else [])

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Requisitos'
        sheet.append(headers)
        for cell in sheet[1]:
            cell.alignment = Alignment(wrap_text=True, vertical='center')

        for req in self._requirements:
            sheet.append([req.id, req.text, req.comment, req.status])
            row = sheet.max_row
            for column in ('A', 'B', 'C', 'D'):
                sheet[f'{column}{row}'].alignment = wrap_top

            if image_resolver and req.images:
                self._embed_image(sheet, row, image_resolver, req.images[0])

        sheet.column_dimensions['A'].width = 38
        sheet.column_dimensions['B'].width = 60
        sheet.column_dimensions['C'].width = 45
        sheet.column_dimensions['D'].width = 24
        if image_resolver:
            sheet.column_dimensions[IMAGE_COLUMN].width = IMAGE_MAX_SIZE / 7
        sheet.freeze_panes = 'A2'

        if option_names:
            options_sheet = workbook.create_sheet('Opciones')
            options_sheet.append([self.STATUS_ATTR])
            for name in option_names:
                options_sheet.append([name])
            options_sheet.sheet_state = 'hidden'

            last_row = len(self._requirements) + 1
            last_option_row = len(option_names) + 1
            validation = DataValidation(
                type='list',
                formula1=f"=Opciones!$A$2:$A${last_option_row}",
                allow_blank=True,
            )
            sheet.add_data_validation(validation)
            validation.add(f'D2:D{last_row}')

        workbook.save(destination)

    def _embed_image(self, sheet, row, image_resolver, image_ref):
        try:
            source = image_resolver(image_ref)
            if source is None:
                return
            if not hasattr(source, 'read') and not Path(source).exists():
                return
            buffer, (width, height) = self._scaled_png(source)
            xl_image = XLImage(buffer)
            xl_image.width, xl_image.height = width, height
            sheet.add_image(xl_image, f'{IMAGE_COLUMN}{row}')
            current_height = sheet.row_dimensions[row].height or 0
            sheet.row_dimensions[row].height = max(current_height, height * 0.75)
        except Exception:
            # corrupted image, format not supported by PIL, etc.: it is
            # skipped instead of breaking the whole export.
            pass

    def update_from_excel(self, source) -> int:
        """Reads an .xlsx (with the columns from to_excel) from `source`
        (path or file-like object) and updates only the comment and
        status of each row, locating the requirement by its ID
        (column A). Returns the number of rows applied."""
        workbook = load_workbook(source, data_only=True)
        sheet = workbook['Requisitos'] if 'Requisitos' in workbook.sheetnames else workbook.active

        updates = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            identifier = str(row[0]).strip()
            comment = row[2] if len(row) > 2 and row[2] is not None else ''
            status = row[3] if len(row) > 3 and row[3] is not None else ''
            updates.append((identifier, str(comment), str(status)))

        return self.update_many(updates)
